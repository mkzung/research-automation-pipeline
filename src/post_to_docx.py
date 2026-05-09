"""Generic DOCX assembly for research-automation pipelines.

Takes a JSON post spec → emits a .docx with title, body paragraphs, embedded
images (charts), optional footer with metadata.

Spec example:

    {
        "title": "Q4 revenue analysis",
        "body": "Long-form narrative text. Paragraphs separated by blank lines.\\n\\n...",
        "charts": ["/output/chart_1.png", "/output/chart_2.png"],
        "footer": "study_id=abc123 · 2026-05-09",
        "tags": ["#datascience", "#analytics"]
    }
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor


NAVY = RGBColor(0x1E, 0x3A, 0x5F)
GRAY_FOOTER = RGBColor(0x9A, 0x9A, 0x9A)
GRAY_TAG = RGBColor(0x6B, 0x6B, 0x6B)


def build_docx(spec: dict, output_path: str | Path) -> Path:
    doc = Document()

    # Default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Title (heading 1)
    if title := spec.get("title"):
        h = doc.add_heading(title, level=1)
        h.runs[0].font.color.rgb = NAVY
        doc.add_paragraph()

    # Body paragraphs (split on double-newline)
    if body := spec.get("body"):
        for para in body.strip().split("\n\n"):
            p = doc.add_paragraph(para.strip())
            p.paragraph_format.space_after = Pt(10)

    # Embedded charts
    for chart_path in spec.get("charts", []):
        if not Path(chart_path).exists():
            continue
        doc.add_paragraph()
        doc.add_picture(chart_path, width=Inches(6.5))

    # Tags / hashtags
    if tags := spec.get("tags"):
        doc.add_paragraph()
        tag_p = doc.add_paragraph()
        run = tag_p.add_run(" ".join(tags))
        run.font.size = Pt(10)
        run.font.color.rgb = GRAY_TAG

    # Footer with metadata
    if footer := spec.get("footer"):
        doc.add_paragraph()
        f = doc.add_paragraph()
        run = f.add_run(footer)
        run.font.size = Pt(8)
        run.font.italic = True
        run.font.color.rgb = GRAY_FOOTER

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def main() -> None:
    p = argparse.ArgumentParser(description="Build a DOCX from a JSON post spec.")
    p.add_argument("--spec", required=True, help="path to spec JSON or '-' for stdin")
    p.add_argument("--output", required=True, help="output .docx path")
    args = p.parse_args()

    if args.spec == "-":
        spec = json.load(sys.stdin)
    else:
        spec = json.loads(Path(args.spec).read_text())

    out = build_docx(spec, args.output)
    print(json.dumps({"ok": True, "output": str(out)}))


if __name__ == "__main__":
    main()
